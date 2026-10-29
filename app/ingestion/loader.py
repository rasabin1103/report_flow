import csv
import io
import json
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

import pandas as pd
from openpyxl import load_workbook

from app.core.settings import ReportError, Settings
from app.schemas.config import Dataset


@dataclass
class Loaded:
    frame: pd.DataFrame
    filename: str
    file_type: str
    sheets: list[str]
    selected_sheet: str | None
    warnings: list[str]


def infer(frame: pd.DataFrame) -> pd.DataFrame:
    for col in frame:
        series = frame[col]
        present = series.dropna()
        if present.empty:
            continue
        text = present.astype(str).str.strip()
        if text.str.lower().isin(["true", "false"]).all():
            frame[col] = series.map(lambda v: None if pd.isna(v) else str(v).lower() == "true").astype(
                "boolean"
            )
        elif pd.to_numeric(present, errors="coerce").notna().all():
            frame[col] = pd.to_numeric(series, errors="coerce")
        elif text.str.match(r"^\d{4}-\d{2}-\d{2}(?:[ T].*)?$").all():
            parsed = pd.to_datetime(series, errors="coerce", format="mixed", utc=True)
            if parsed.notna().sum() == len(present):
                frame[col] = parsed
    return frame


def load(data: bytes, filename: str, options: Dataset, settings: Settings) -> Loaded:
    if not data:
        raise ReportError("The dataset is empty")
    if len(data) > settings.max_file_bytes:
        raise ReportError("File size limit exceeded", 413)
    suffix = PurePath(filename).suffix.lower()
    sheets: list[str] = []
    selected = None
    warnings: list[str] = []
    rows: list[list] = []
    try:
        if suffix == ".csv":
            text = data.decode(options.encoding)
            delimiter = options.delimiter
            if delimiter is None:
                try:
                    delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
                except csv.Error:
                    delimiter = ","
            reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            header = next(reader)
            if not header or len(header) > settings.max_columns:
                raise ReportError("Column limit exceeded or empty header", 413)
            for row in reader:
                if not row:
                    continue
                if len(row) != len(header):
                    raise ReportError("Malformed CSV: inconsistent row width")
                rows.append([None if v == "" else v for v in row])
                if len(rows) > settings.max_rows:
                    raise ReportError("Row limit exceeded", 413)
        elif suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                infos = archive.infolist()
                if len(infos) > 2000 or sum(i.file_size for i in infos) > settings.max_uncompressed_bytes:
                    raise ReportError("Excel expansion limit exceeded", 413)
                if any("vbaproject" in i.filename.lower() for i in infos):
                    raise ReportError("Macro-enabled workbooks are unsupported")
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                sheets = book.sheetnames
                selected = options.sheet_name or sheets[0]
                if selected not in sheets:
                    raise ReportError("Selected worksheet does not exist")
                sheet = book[selected]
                if (sheet.max_column or 0) > settings.max_columns:
                    raise ReportError("Column limit exceeded", 413)
                if (sheet.max_row or 0) > settings.max_rows + 1:
                    raise ReportError("Row limit exceeded", 413)
                sheet.reset_dimensions()  # Scan physical rows even when producer dimensions understate them.
                iterator = sheet.iter_rows(values_only=True)
                header = list(next(iterator))
                for row in iterator:
                    if len(row) > len(header):
                        raise ReportError("Worksheet row contains values beyond the header")
                    rows.append(list(row) + [None] * (len(header) - len(row)))
                    if len(row) > settings.max_columns or len(rows) > settings.max_rows:
                        raise ReportError("Dataset dimensions exceed limits", 413)
            finally:
                book.close()
            warnings.append(
                "Only stored formula results are loaded; formulas without cached values appear as null."
            )
        else:
            raise ReportError("Unsupported format; use .csv or .xlsx")
        names = [str(v).strip() if v is not None else "" for v in header]
        if not names or any(not n for n in names):
            raise ReportError("Column names must be nonempty")
        if len(names) != len(set(names)):
            raise ReportError("Duplicate column names")
        if len(names) > settings.max_columns:
            raise ReportError("Column limit exceeded", 413)
        frame = infer(pd.DataFrame(rows, columns=names))
        # Bound cell text before it reaches chart or paragraph layout.
        if any(frame[c].dropna().astype(str).str.len().gt(2000).any() for c in frame):
            raise ReportError("Cell text exceeds 2000 characters")
        return Loaded(
            frame, PurePath(filename.replace("\\", "/")).name, suffix[1:], sheets, selected, warnings
        )
    except ReportError:
        raise
    except Exception:  # noqa: BLE001 - sanitize untrusted parser errors at the boundary
        raise ReportError(
            "Cannot read dataset: check format, encoding, headers and workbook integrity"
        ) from None


def column_type(series: pd.Series) -> str:
    if pd.api.types.is_bool_dtype(series):
        return "boolean"
    if pd.api.types.is_numeric_dtype(series):
        return "number"
    if pd.api.types.is_datetime64_any_dtype(series):
        return "date"
    return "string"


def inspect_dataset(loaded: Loaded) -> dict:
    frame = loaded.frame
    return {
        "file_type": loaded.file_type,
        "sheet_names": loaded.sheets,
        "selected_sheet": loaded.selected_sheet,
        "total_rows": len(frame),
        "total_columns": len(frame.columns),
        "column_names": list(frame.columns),
        "columns": [
            {
                "name": c,
                "type": column_type(frame[c]),
                "null_count": int(frame[c].isna().sum()),
                "unique_count": int(frame[c].nunique()),
            }
            for c in frame
        ],
        "sample": json.loads(frame.head(5).to_json(orient="records", date_format="iso")),
        "warnings": loaded.warnings,
    }
